"""Builds the free Lowest-Week Calculator (xlsx, opens cleanly in Google Sheets)."""
import sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation

SAGE, CREAM, TERRA, CHAR, GREEN, YELLOW = "8FA98B", "F4EFE6", "C0563A", "2B2B2B", "DCE8D5", "FFF1C9"
out = sys.argv[1] if len(sys.argv) > 1 else "Lowest-Week-Calculator.xlsx"

wb = Workbook()
ws = wb.active
ws.title = "Calculator"
ws.sheet_view.showGridLines = False
for col, w in {"A": 3, "B": 38, "C": 18, "D": 46}.items():
    ws.column_dimensions[col].width = w
fill = lambda c: PatternFill("solid", fgColor=c)
thin = Side(style="thin", color="D8D2C6")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
for r in range(1, 40):
    for c in "ABCDE":
        ws[f"{c}{r}"].fill = fill(CREAM)

ws["B2"] = "Lowest-Week Calculator"
ws["B2"].font = Font(name="Georgia", size=22, bold=True, color=CHAR)
ws["B3"] = "Find the weekly number you can always count on, even when every paycheck is different."
ws["B3"].font = Font(size=11, italic=True, color=CHAR)
ws["B5"] = "1. Type what you earned each week (rough is fine). 2. Read your safe number in the green box."
ws["B5"].font = Font(size=10, color=CHAR)

ws["B7"], ws["C7"] = "Week", "Take-home pay ($)"
for c in ("B7", "C7"):
    ws[c].font = Font(bold=True, color="FFFFFF")
    ws[c].fill = fill(SAGE)
    ws[c].border = box
examples = [620, 380, 540, 290, 710, 450, 330, 600]
for i in range(8):
    r = 8 + i
    ws[f"B{r}"] = f"Week {i + 1}" + (" (most recent)" if i == 7 else "")
    ws[f"C{r}"] = examples[i]
    ws[f"C{r}"].number_format = '"$"#,##0'
    ws[f"C{r}"].fill = fill("FFFFFF")
    for c in "BC":
        ws[f"{c}{r}"].border = box
ws["D8"] = "Example numbers. Replace them with yours."
ws["D8"].font = Font(size=9, italic=True, color=TERRA)

dv = DataValidation(type="decimal", operator="between", formula1="0", formula2="100000", allow_blank=True,
                    showErrorMessage=True, errorTitle="Numbers only", error="Type a dollar amount, like 450.")
ws.add_data_validation(dv)
dv.add("C8:C15")

R = "C8:C15"
rows = [
    (18, "Your average week", f'=IF(COUNT({R})<3,"",AVERAGE({R}))', "What most budgets use. It hides your slow weeks."),
    (19, "Weeks below your average", f'=IF(COUNT({R})<3,"",COUNTIF({R},"<"&C18))', "These are the weeks an average-based budget comes up short."),
    (21, "YOUR SAFE WEEKLY NUMBER", f'=IF(COUNT({R})<3,"Type at least 3 weeks",(SMALL({R},1)+SMALL({R},2)+SMALL({R},3))/3)',
     "The average of your 3 lowest weeks. Plan your must-pay bills around this."),
    (22, "Safe number per month", f'=IF(ISNUMBER(C21),C21*52/12,"")', "Weekly safe number x 52 weeks / 12 months."),
    (23, "Extra in an average week", f'=IF(ISNUMBER(C21),C18-C21,"")', "In a typical week, this much can go to your cushion for slow weeks."),
]
for r, label, f, note in rows:
    ws[f"B{r}"], ws[f"C{r}"], ws[f"D{r}"] = label, f, note
    ws[f"C{r}"].number_format = '"$"#,##0' if r != 19 else "0"
    ws[f"D{r}"].font = Font(size=9, color=CHAR)
    ws[f"D{r}"].alignment = Alignment(wrap_text=True, vertical="center")
    for c in "BC":
        ws[f"{c}{r}"].border = box
        ws[f"{c}{r}"].fill = fill("FFFFFF")
    ws[f"B{r}"].font = Font(bold=True, color=CHAR)
for c in ("B21", "C21"):
    ws[c].fill = fill(GREEN)
    ws[c].font = Font(bold=True, size=14, color=CHAR)
ws.row_dimensions[21].height = 30
ws["C19"].font = Font(bold=True, color=TERRA)
ws.conditional_formatting.add("C19", FormulaRule(formula=["AND(ISNUMBER(C19),C19>0)"], fill=fill(YELLOW)))

ws["B26"] = "Want the sheet that tells you what to do with this number every payday?"
ws["B26"].font = Font(bold=True, color=TERRA)
ws["B27"] = ("The Uneven Paycheck Kit splits every paycheck into bills set-aside, tax jar, safe to spend "
             "and cushion, with a bills calendar and a rent-week plan.")
ws["B27"].alignment = Alignment(wrap_text=True)
ws.merge_cells("B27:D27")
ws.row_dimensions[27].height = 32
ws["B29"] = "Made by Nora, an AI creator. Template built and tested by a real human. Not tax, debt or investing advice."
ws["B29"].font = Font(size=8, italic=True, color="7A7468")
ws.freeze_panes = "A7"
wb.save(out)
print("saved", out)
